# -*- coding: utf-8 -*-
"""
Lecture d'un bilan ODF au format Word (.docx) et extraction structurée des
éléments (identité, analyse de Steiner, analyse d'espace/DDM, synthèse
diagnostique, motif, objectifs & moyens) pour les reporter dans le dossier
patient de l'application.

Le lecteur repère les tableaux par leur CONTENU (pas par leur position), pour
rester fiable même si la mise en page varie un peu d'un bilan à l'autre.
Calé sur le format des bilans fournis (rubriques : ANALYSE STEINER, analyse
d'espace +/−, SYNTHESE DIAGNOSTIQUE en grille Squelettique/Alvéolaire/Occlusal/
Cutané × Sagittal/Transversal/Vertical, STRATEGIES ET MOYENS).
"""
import re
import unicodedata


def _strip_accents(s):
    s = unicodedata.normalize("NFKD", s or "")
    return "".join(c for c in s if not unicodedata.combining(c))


def _norm(s):
    """minuscule, sans accents, espaces normalisés (espaces insécables inclus)."""
    s = (s or "").replace("\xa0", " ").replace(" ", " ")
    s = _strip_accents(s).lower()
    return re.sub(r"\s+", " ", s).strip()


def _canon(s):
    """clé compacte : lettres/chiffres uniquement (pour comparer des libellés)."""
    return re.sub(r"[^a-z0-9]", "", _norm(s))


def _num(s):
    """premier nombre (virgule ou point) trouvé, sinon None. '-' -> None."""
    if s is None:
        return None
    s = str(s).replace("\xa0", " ").replace(" ", " ").strip().replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return None
    try:
        return float(m.group(0))
    except Exception:
        return None


# ---- Steiner : clé compacte du libellé Word -> clé interne de l'application ----
STEINER_CANON = {
    "sna": "SNA (°)",
    "snb": "SNB (°)",
    "anb": "ANB (°)",
    "aobomm": "Wits (mm)",          # AoBo = Wits
    "inamm": "U1-NA (mm)",
    "ina": "U1-NA (°)",
    "inbmm": "L1-NB (mm)",
    "inb": "L1-NB (°)",
    "pognbmm": "Pog-NB (mm)",
    "ii": "Angle interincisif (°)",
    "osn": "Plan occlusal/SN (°)",
    "gognsn": "Plan mandibulaire (Go-Gn/SN) (°)",
}

# ---- Analyse d'espace : clé compacte du libellé Word -> clé interne ----
ESPACE_CANON = {
    "encombrement": "encombrement",
    "repositionnementincisive": "repositionnement_incisive",
    "courbespee": "courbe_spee",
    "repositionnement6": "repositionnement_6",
    "espacederivemesiale": "espace_derive_mesiale",
    "extractionstripping": "extraction_stripping",
    "individualisation": "individualisation",
}

# ---- Synthèse : entêtes de colonnes / lignes de la grille ----
_COL_CANON = {"squelettique": "sq", "alveolaire": "al", "occlusal": "oc", "cutane": "cu"}


def _row_key(label):
    n = _norm(label)
    if "transvers" in n:
        return "tr"
    if "vertical" in n:
        return "ve"
    if any(w in n for w in ("anter", "poster", "sagittal", "a-p", "ap")):
        return "ap"
    return None


def _cell_text(cell):
    """texte d'une cellule en préservant les retours à la ligne (paragraphes)."""
    parts = [p.text for p in cell.paragraphs]
    txt = "\n".join(parts)
    return txt.replace("\xa0", " ").replace(" ", " ")


def _rows_as_text(table):
    out = []
    for row in table.rows:
        out.append([_cell_text(c).strip() for c in row.cells])
    return out


def _table_blob(table):
    return _norm(" ".join(_cell_text(c) for row in table.rows for c in row.cells))


def _clean_multiline(txt):
    lines = [ln.strip() for ln in (txt or "").replace("\r", "\n").split("\n")]
    lines = [ln for ln in lines if ln]
    return "\n".join(lines)


# ---- Chevrons : libellé de la colonne (row du bas) -> clé de chevron de l'app ----
CHEVRON_CANON = {"probleme": "probleme", "solution": "solution", "individualisation": "individualisation"}


def _fmt_val(v):
    return ("%g" % v)


def parse_chevrons(tables):
    """Décode le tableau des chevrons de Steiner.
    Chaque chevron « > » (ouvert à gauche) occupe 2 colonnes :
      - colonne EXTERNE (gauche)  : ANB (haut) / Pog-NB (bas)
      - colonne INTERNE (droite)  : I-NA mm (haut) / I-NB mm (bas)
    La dernière ligne porte les libellés (Problème / Solution / Individualisation).
    Retour : {clé_chevron: {'anb':..,'ina':..,'inb':..,'pog':..}} (valeurs = str)."""
    out = {}
    target = None
    for t in tables:
        rows = _rows_as_text(t)
        if not rows:
            continue
        last = [_canon(c) for c in rows[-1]]
        if any(c in CHEVRON_CANON for c in last):
            target = rows
            break
    if not target:
        return out
    label_row = target[-1]
    ncols = max(len(r) for r in target)
    nlab = len(label_row)
    # regroupe les colonnes par libellé de chevron
    regions = {}  # clé_chevron -> liste d'indices de colonnes
    for i in range(nlab):
        key = CHEVRON_CANON.get(_canon(label_row[i]))
        if key:
            regions.setdefault(key, []).append(i)

    def col_values(col):
        """[(row_index, valeur)] non vides de la colonne, hors ligne de libellés."""
        vals = []
        for ri, r in enumerate(target[:-1]):
            if col < len(r):
                v = _num(r[col])
                if v is not None:
                    vals.append((ri, v))
        vals.sort()
        return vals

    for key, cols in regions.items():
        outer = min(cols)
        inner = outer + 1
        ov = col_values(outer)
        iv = col_values(inner) if inner < ncols else []
        d = {}
        mid = (len(target) - 1) / 2.0
        if len(ov) >= 2:
            d["anb"] = _fmt_val(ov[0][1]); d["pog"] = _fmt_val(ov[-1][1])
        elif len(ov) == 1:
            (r, v) = ov[0]
            d["anb" if r < mid else "pog"] = _fmt_val(v)
        if len(iv) >= 2:
            d["ina"] = _fmt_val(iv[0][1]); d["inb"] = _fmt_val(iv[-1][1])
        elif len(iv) == 1:
            (r, v) = iv[0]
            d["ina" if r < mid else "inb"] = _fmt_val(v)
        if d:
            out[key] = d
    return out


def parse_word_bilan(path):
    """Retourne un dict :
    {patient:{nom,sexe,age,dob}, steiner:{clé:val}, espace:{clé:{p|m}},
     synthese:{clé:texte}, motif:str, objectifs:str, moyens:str,
     notes:[str]}  (notes = éléments repérés mais sans emplacement dédié)
    """
    import docx  # python-docx
    d = docx.Document(path)
    res = {"patient": {}, "steiner": {}, "espace": {}, "synthese": {}, "chevrons": {},
           "motif": "", "objectifs": "", "moyens": "", "notes": []}
    tables = d.tables

    # --- Identité : 1er tableau contenant nom + (Homme/Femme ou une date) ---
    for t in tables:
        rows = _rows_as_text(t)
        if not rows or not rows[0]:
            continue
        cells = rows[0]
        blob = _norm(" ".join(cells))
        if cells[0] and (("homme" in blob or "femme" in blob) or re.search(r"\d{1,2}/\d{1,2}/\d{4}", blob)):
            res["patient"]["nom"] = cells[0].strip()
            meta = " ".join(cells[1:])
            mn = _norm(meta)
            if "homme" in mn or "garcon" in mn:
                res["patient"]["sexe"] = "masculin"
            elif "femme" in mn or "fille" in mn:
                res["patient"]["sexe"] = "féminin"
            ma = re.search(r"(\d+)\s*ans", mn)
            if ma:
                res["patient"]["age"] = "%s ans" % ma.group(1)
            md = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", meta)
            if md:
                dd, mm, yy = md.group(1), md.group(2), md.group(3)
                res["patient"]["dob"] = "%s-%02d-%02d" % (yy, int(mm), int(dd))
            break

    # --- Analyse de Steiner : tableau avec des libellés type SNA/ANB/I-NA ---
    for t in tables:
        rows = _rows_as_text(t)
        canon_labels = [_canon(r[0]) for r in rows if r]
        if not any(c in STEINER_CANON for c in canon_labels):
            continue
        for r in rows:
            if len(r) < 3:
                continue
            key = STEINER_CANON.get(_canon(r[0]))
            if not key:
                continue
            # colonne 1 = norme (moyenne), colonne 2 = valeur du patient
            val = _num(r[2]) if len(r) > 2 else None
            if val is None and len(r) > 3:
                val = _num(r[3])
            if val is not None:
                res["steiner"][key] = val
        break

    # --- Analyse d'espace / DDM : tableau dont l'entête contient + et − ---
    for t in tables:
        rows = _rows_as_text(t)
        if not rows:
            continue
        head = [_norm(c) for c in rows[0]]
        if not ("+" in [h.strip() for h in head] and "-" in [h.strip() for h in head]):
            # certaines cellules peuvent contenir le signe entouré d'espaces
            if not (any(h.strip() == "+" for h in head) and any(h.strip() in ("-", "−") for h in head)):
                continue
        # indices des colonnes + et −
        try:
            ip = next(i for i, h in enumerate(rows[0]) if _norm(h).strip() == "+")
            im = next(i for i, h in enumerate(rows[0]) if _norm(h).strip() in ("-", "−"))
        except StopIteration:
            ip, im = 1, 2
        for r in rows[1:]:
            key = ESPACE_CANON.get(_canon(r[0]))
            if not key:
                continue
            cell = {}
            p = _num(r[ip]) if len(r) > ip else None
            m = _num(r[im]) if len(r) > im else None
            if p is not None:
                cell["p"] = round(p, 2)
            if m is not None:
                cell["m"] = round(m, 2)
            if cell:
                res["espace"][key] = cell
        break

    # --- Synthèse diagnostique : grille Squelettique/Alvéolaire/Occlusal/Cutané ---
    for t in tables:
        rows = _rows_as_text(t)
        if not rows:
            continue
        head = [_canon(c) for c in rows[0]]
        if not any(c in _COL_CANON for c in head):
            continue
        colmap = {}  # index de colonne -> clé (sq/al/oc/cu)
        for i, c in enumerate(head):
            if c in _COL_CANON:
                colmap[i] = _COL_CANON[c]
        for r in rows[1:]:
            rk = _row_key(r[0])
            if not rk:
                continue
            for i, ck in colmap.items():
                if i < len(r):
                    txt = r[i].strip()
                    txt = re.sub(r"\s*\n\s*", " / ", txt).strip(" /")
                    if txt and txt not in ("-", "—"):
                        res["synthese"]["%s_%s" % (ck, rk)] = txt
        break

    # --- Motif + objectifs/moyens : parcours des cellules ---
    for t in tables:
        for row in t.rows:
            for c in row.cells:
                txt = _cell_text(c)
                n = _norm(txt)
                if not res["motif"] and "motif de consultation" in n:
                    m = re.search(r"motif de consultation\s*:?\s*(.*?)(?:\bhbd\b|\bparodonte\b|\bpano\b|$)",
                                  txt.replace("\n", " "), re.I | re.S)
                    if m:
                        res["motif"] = m.group(1).strip(" :/-\t")
                if ("objectifs" in n) and ("moyens" in n):
                    mo = re.search(r"objectifs\s*:?(.*?)moyens\s*:?", txt, re.I | re.S)
                    mm = re.search(r"moyens\s*:?(.*)$", txt, re.I | re.S)
                    if mo:
                        res["objectifs"] = _clean_multiline(mo.group(1))
                    if mm:
                        res["moyens"] = _clean_multiline(mm.group(1))

    # --- Chevrons ---
    try:
        res["chevrons"] = parse_chevrons(tables)
    except Exception:
        res["chevrons"] = {}

    return res
